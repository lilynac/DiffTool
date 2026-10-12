"""Read-only comparison engine; independent of the desktop UI."""
from dataclasses import dataclass
from difflib import SequenceMatcher
from itertools import zip_longest
from pathlib import Path
import csv
import io
from .alignment import alignment_keys, align_rows, mapped_formula

EXCEL = {'.xlsx', '.xlsm'}
MAX_CELLS = 2_000_000

@dataclass(frozen=True)
class Options:
    ignore_space: bool = False
    ignore_case: bool = False
    cached_values: bool = False
    encoding: str = 'auto'
    align_excel: bool = True

@dataclass
class Row:
    left_number: int | None
    right_number: int | None
    left: tuple
    right: tuple
    changed: tuple[bool, ...]
    status: str

@dataclass
class Section:
    name: str
    rows: list[Row]
    columns: int
    presence: str = '両側'
    is_text: bool = False

    @property
    def differences(self):
        return sum(r.status != 'equal' for r in self.rows)


def normalized(value, options):
    if isinstance(value, str):
        if options.ignore_space:
            value = ''.join(value.split())
        if options.ignore_case:
            value = value.casefold()
    # Preserve distinctions such as numeric 1 versus text "1" and bool True.
    return (type(value).__name__, value)


def compare_rows(left, right, options=Options(), align=True, excel=False, sheet_name=''):
    norm = lambda value: normalized(value, options)
    if align and excel:
        a, b = alignment_keys(left, norm), alignment_keys(right, norm)
        pairs = align_rows(a, b)
    elif align:
        a = [tuple(norm(v) for v in row) for row in left]
        b = [tuple(norm(v) for v in row) for row in right]
        pairs = [(li, ri) for _, i, j, k, l in SequenceMatcher(None, a, b, autojunk=False).get_opcodes()
                 for li, ri in zip_longest(range(i, j), range(k, l))]
    else:
        pairs = list(zip_longest(range(len(left)), range(len(right))))
    row_map = {ri+1: li+1 for li, ri in pairs if li is not None and ri is not None}

    def different(lv, rv):
        if norm(lv) == norm(rv):
            return False
        if excel and align:
            return norm(lv) != norm(mapped_formula(rv, row_map, sheet_name))
        return True

    result = []
    for li, ri in pairs:
        lv = left[li] if li is not None else ()
        rv = right[ri] if ri is not None else ()
        width = max(len(lv), len(rv), 1)
        changed = tuple(li is None or ri is None or
                        different(lv[c] if c < len(lv) else None,
                                  rv[c] if c < len(rv) else None)
                        for c in range(width))
        status = 'added' if li is None else 'removed' if ri is None else 'changed' if any(changed) else 'equal'
        result.append(Row(None if li is None else li+1, None if ri is None else ri+1,
                          lv, rv, changed, status))
    return result


def decode(path, encoding):
    raw = Path(path).read_bytes()
    if len(raw) > 50_000_000:
        raise ValueError('テキストは50MB以下のファイルを指定してください。')
    if encoding != 'auto':
        return raw.decode(encoding)
    candidates = ['utf-16'] if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else ['utf-8-sig', 'cp932']
    for candidate in candidates:
        try:
            return raw.decode(candidate)
        except UnicodeDecodeError:
            pass
    raise ValueError('文字コードを判定できません。文字コードを指定してください。')


def read_text(path, options, structured):
    text = decode(path, options.encoding)
    if not structured:
        # Retain line terminators so missing final newline is visible.
        return [(line,) for line in text.splitlines(keepends=True)]
    delimiter = '\t' if Path(path).suffix.lower() == '.tsv' else ','
    return [tuple(row) for row in csv.reader(io.StringIO(text, newline=''), delimiter=delimiter)]


def read_excel(path, options):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=options.cached_values)
    try:
        result = {}
        total = 0
        for sheet in book.worksheets:
            size = (sheet.max_row or 0) * (sheet.max_column or 0)
            total += size
            if total > MAX_CELLS:
                raise ValueError('Excelは1ファイル200万セルまでです。不要な行・列の書式を削除してください。')
            rows = []
            for values in sheet.iter_rows(values_only=True):
                values = list(values)
                while values and values[-1] is None:
                    values.pop()
                rows.append(tuple(values))
            while rows and not rows[-1]:
                rows.pop()
            result[sheet.title] = rows
        return result
    finally:
        book.close()


def compare_files(left, right, options=Options()):
    lp, rp = Path(left), Path(right)
    if '.xls' in (lp.suffix.lower(), rp.suffix.lower()):
        raise ValueError('旧形式 .xls は未対応です。.xlsx に変換してください。')
    excel = lp.suffix.lower() in EXCEL
    if excel != (rp.suffix.lower() in EXCEL):
        raise ValueError('Excel同士、またはテキスト同士を選択してください。')
    if excel:
        a, b = read_excel(lp, options), read_excel(rp, options)
        sections = []
        for name in dict.fromkeys([*a, *b]):
            rows = compare_rows(a.get(name, []), b.get(name, []), options, align=options.align_excel, excel=True, sheet_name=name)
            presence = '右のみ' if name not in a else '左のみ' if name not in b else '両側'
            sections.append(Section(name, rows, max([len(r.left) for r in rows] + [len(r.right) for r in rows] + [1]), presence))
        return sections
    structured = lp.suffix.lower() in {'.csv', '.tsv'} and rp.suffix.lower() in {'.csv', '.tsv'}
    rows = compare_rows(read_text(lp, options, structured), read_text(rp, options, structured), options)
    return [Section('CSV / TSV' if structured else 'テキスト', rows,
                    max([len(r.left) for r in rows] + [len(r.right) for r in rows] + [1]), is_text=not structured)]


def export_report(path, sections):
    with open(path, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['シート', '状態', '左行', '右行', '列', '左の値', '右の値'])
        for section in sections:
            if section.presence != '両側':
                writer.writerow([section.name, section.presence, '', '', '', '', ''])
            for row in section.rows:
                for c, changed in enumerate(row.changed):
                    if changed:
                        writer.writerow([section.name, row.status, row.left_number, row.right_number, c+1,
                                         row.left[c] if c < len(row.left) else '', row.right[c] if c < len(row.right) else ''])
