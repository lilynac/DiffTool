import csv
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from openpyxl import Workbook
from compare_tool.core import Options, compare_rows, compare_files, export_report


class ComparisonTests(unittest.TestCase):
    def test_insert_alignment(self):
        rows = compare_rows([('a',), ('b',)], [('a',), ('new',), ('b',)])
        self.assertEqual([r.status for r in rows], ['equal', 'added', 'equal'])
        self.assertEqual(rows[-1].left_number, 2)
        self.assertEqual(rows[-1].right_number, 3)

    def test_options_and_types(self):
        self.assertEqual(compare_rows([(' A B ',)], [('ab',)], Options(True, True))[0].status, 'equal')
        self.assertEqual(compare_rows([(1,)], [('1',)])[0].status, 'changed')
        self.assertEqual(compare_rows([(True,)], [(1,)])[0].status, 'changed')

    def test_excel_cells_and_missing_sheet(self):
        with TemporaryDirectory() as directory:
            a, b = Path(directory)/'a.xlsx', Path(directory)/'b.xlsx'
            book = Workbook()
            book.active.title = '売上'
            book.active.append(['商品', '数量'])
            book.active.append(['りんご', 1])
            book.create_sheet('左だけ')
            book.save(a)
            book['売上']['B2'] = 2
            del book['左だけ']
            book.create_sheet('右だけ')
            book.save(b)
            sections = compare_files(a, b)
            self.assertEqual([s.name for s in sections], ['売上', '左だけ', '右だけ'])
            self.assertEqual(sections[0].rows[1].changed, (False, True))
            self.assertEqual(sections[1].presence, '左のみ')
            self.assertEqual(sections[2].presence, '右のみ')
            report = Path(directory)/'report.csv'
            export_report(report, sections)
            with report.open(encoding='utf-8-sig', newline='') as f:
                records = list(csv.reader(f))
            self.assertEqual(records[1][2:5], ['2', '2', '2'])
            self.assertEqual(len(records), 4)

    def test_csv_quotes_and_japanese_encoding(self):
        with TemporaryDirectory() as directory:
            a, b = Path(directory)/'a.csv', Path(directory)/'b.csv'
            a.write_bytes('名前,備考\r\n太郎,"a,b"\r\n'.encode('cp932'))
            b.write_text('名前,備考\n太郎,"a,c"\n', encoding='utf-8-sig')
            section = compare_files(a, b)[0]
            self.assertEqual(section.rows[0].status, 'equal')
            self.assertEqual(section.rows[1].changed, (False, True))

    def test_line_endings_and_empty_file(self):
        with TemporaryDirectory() as directory:
            a, b = Path(directory)/'a.txt', Path(directory)/'b.txt'
            a.write_bytes(b'a\r\n')
            b.write_bytes(b'a\n')
            self.assertEqual(compare_files(a, b)[0].differences, 1)
            a.write_text('')
            b.write_text('')
            self.assertEqual(compare_files(a, b)[0].rows, [])

    def test_excel_fixed_coordinates(self):
        rows = compare_rows([('a',), ('b',)], [('x',), ('a',), ('b',)], align=False)
        self.assertEqual([r.status for r in rows], ['changed', 'changed', 'added'])

if __name__ == '__main__':
    unittest.main()
